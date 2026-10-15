"""简历文件解析：PDF / DOCX → 纯文本。

对纯图片 PDF 必须**明确报错**，而不是返回空字符串让 LLM 去猜。
"""

from __future__ import annotations

import io
import logging
import re

from app.core.errors import DocumentParseError, InvalidInputError

logger = logging.getLogger("app.parser")

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")

# PDF 里可能的 MIME 类型
_PDF_MIME = ("application/pdf",)
_DOCX_MIME = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
)
_TEXT_MIME = ("text/plain", "text/markdown", "")

_MULTI_BLANK_RE = re.compile(r"\n{3,}")
_TRAILING_SPACE_RE = re.compile(r"[ \t]+\n")


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _TRAILING_SPACE_RE.sub("\n", text)
    text = _MULTI_BLANK_RE.sub("\n\n", text)
    return text.strip()


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    # 优先使用 pdfplumber（版面还原更好），未安装则回退 pypdf
    try:
        import pdfplumber  # type: ignore

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [(p.extract_text() or "") for p in pdf.pages]
        text = "\n".join(pages)
        if text.strip():
            return text
        logger.info("pdfplumber 未提取到文字，回退 pypdf")
    except ImportError:
        pass
    except Exception as exc:
        logger.warning("pdfplumber 解析失败(%s)，回退 pypdf", type(exc).__name__)

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise DocumentParseError("该 PDF 已加密，请先解除密码保护后再上传") from exc
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except PdfReadError as exc:
        raise DocumentParseError("PDF 文件已损坏或格式不受支持") from exc


def _extract_docx(data: bytes) -> str:
    import docx  # python-docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentParseError("Word 文件已损坏或格式不受支持（仅支持 .docx）") from exc

    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]

    # 表格里常有技能清单，别漏掉
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes, *, max_bytes: int) -> str:
    """把上传文件转成纯文本。

    抛出的是 AppError 子类，路由无需再包一层 try/except。
    """
    if not data:
        raise InvalidInputError("上传的文件是空的")
    if len(data) > max_bytes:
        raise InvalidInputError(f"文件超过 {max_bytes // 1024 // 1024}MB 限制")

    lower = filename.lower()
    if lower.endswith(".pdf"):
        text = _extract_pdf(data)
    elif lower.endswith(".docx"):
        text = _extract_docx(data)
    elif lower.endswith((".txt", ".md")):
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise DocumentParseError("文本文件不是 UTF-8 编码，请另存为 UTF-8 后重试") from exc
    else:
        raise InvalidInputError(
            f"不支持的文件类型，仅支持 {'/'.join(SUPPORTED_EXTENSIONS)}"
        )

    text = _clean(text)

    if len(text) < 50:
        raise DocumentParseError(
            "未识别到有效文字。如果你上传的是扫描件或纯图片 PDF，"
            "请改用文字版简历（在 Word 里导出 PDF，或直接上传 .docx）",
            detail={"extracted_length": len(text)},
        )
    return text


def guess_mime(filename: str) -> str:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return _PDF_MIME[0]
    if lower.endswith(".docx"):
        return _DOCX_MIME[0]
    return _TEXT_MIME[0]
